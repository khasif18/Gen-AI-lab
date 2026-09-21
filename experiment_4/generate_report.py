import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_word_document():
    doc = Document()
    
    # 1. Title
    title = doc.add_heading('Laboratory Experiment Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading('Implementation of LLM-Powered Conversational Agent Using Prompt Engineering and API Integration', level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 2. Aim
    doc.add_heading('2. AIM', level=2)
    doc.add_paragraph('To implement an LLM-powered conversational agent (chatbot) utilizing prompt engineering and API integration to serve as a college admission assistant.')
    
    # 3. Objectives
    doc.add_heading('3. OBJECTIVES', level=2)
    doc.add_paragraph('1. To understand the principles and mechanics of LLM-powered conversational agents.', style='List Number')
    doc.add_paragraph('2. To design effective system prompts to constrain and direct the behavior of the LLM.', style='List Number')
    doc.add_paragraph('3. To integrate an external LLM API into a backend web server.', style='List Number')
    doc.add_paragraph('4. To build a modern web-based chatbot interface for student interactions.', style='List Number')
    doc.add_paragraph('5. To accurately answer structured college admission-related queries using conversational context.', style='List Number')
    
    # 4. Theory
    doc.add_heading('4. THEORY', level=2)
    p = doc.add_paragraph()
    p.add_run('Large Language Model (LLM): ').bold = True
    p.add_run('A deep learning algorithm capable of recognizing, summarizing, translating, predicting, and generating human language based on vast amounts of text data.\n')
    p.add_run('Conversational Agent: ').bold = True
    p.add_run('A software program that interprets and responds to user input in a natural, conversational manner. LLMs serve as the "brain" of modern conversational agents.\n')
    p.add_run('Prompt Engineering: ').bold = True
    p.add_run('The process of carefully designing the input text (instructions, context, rules) given to an LLM to elicit a highly specific and accurate output. A "System Prompt" is used to define the agent\'s overarching persona and rules.\n')
    p.add_run('API Integration: ').bold = True
    p.add_run('Application Programming Interface (API) allows the custom backend application to send data (prompts) to the remote LLM servers and receive generated text in response via HTTP protocols.\n')
    p.add_run('Context: ').bold = True
    p.add_run('Because APIs are stateless, the conversation history (previous questions and answers) must be explicitly sent alongside every new query to allow the LLM to understand contextual references (e.g., "What about the second one?").')

    # 5. Case Study
    doc.add_heading('5. CASE STUDY', level=2)
    doc.add_paragraph('Fictional Institution: ABC Institute of Technology\n')
    doc.add_paragraph('The agent is supplied with a structured local JSON knowledge base containing demo/synthetic admission information including:')
    doc.add_paragraph('• Courses: B.Tech (CSE, AI & ML, ECE, Mechanical)', style='List Bullet')
    doc.add_paragraph('• Eligibility: 10+2 passed with Physics & Mathematics; Minimum 60% aggregate.', style='List Bullet')
    doc.add_paragraph('• Fees: Ranging from $4000 to $5500 depending on the course.', style='List Bullet')
    doc.add_paragraph('• Documents: 10th/12th certificates, TC, ID proof, photos.', style='List Bullet')

    # 6. System Architecture
    doc.add_heading('6. SYSTEM ARCHITECTURE', level=2)
    doc.add_paragraph('The system follows a 3-tier architecture with a specialized prompt injection layer:')
    arch = '''Student
   ↓
Chatbot Web Interface (Frontend)
   ↓
User Query (HTTP POST)
   ↓
Prompt Engineering Layer (Flask Backend)
   ↓
College Admission Knowledge (JSON injection) + Conversation Context
   ↓
LLM API (OpenAI)
   ↓
Generated Response
   ↓
Chatbot Interface
   ↓
Student'''
    doc.add_paragraph(arch, style='No Spacing')

    # 7. Prompt Engineering
    doc.add_heading('7. PROMPT ENGINEERING', level=2)
    doc.add_paragraph('The chatbot utilizes a strict system prompt to control behavior:')
    prompt = '''You are the AI Admission Assistant for ABC Institute of Technology.
Your task is to answer student questions about admissions.

Rules:
1. Use ONLY the college information provided below.
2. Do not invent fees, dates, eligibility criteria, courses, or procedures.
3. If the information is unavailable, clearly say so.
4. Maintain context from previous messages.
5. Answer in simple and professional language.
6. Ask for clarification when necessary.'''
    doc.add_paragraph(prompt, style='No Spacing')

    # 8. Algorithm
    doc.add_heading('8. ALGORITHM', level=2)
    doc.add_paragraph('1. Start the Flask server and load the JSON knowledge base.', style='List Number')
    doc.add_paragraph('2. User opens the web UI and types a query.', style='List Number')
    doc.add_paragraph('3. The UI sends the query + recent conversation history to the backend via POST.', style='List Number')
    doc.add_paragraph('4. The backend constructs a message array: [System Prompt] + [History] + [User Query].', style='List Number')
    doc.add_paragraph('5. Backend calls the LLM API and waits for the generated text.', style='List Number')
    doc.add_paragraph('6. The text is returned to the UI.', style='List Number')
    doc.add_paragraph('7. The UI formats the text (lists, bolding) and appends it to the chat box.', style='List Number')

    # 9. Pseudocode
    doc.add_heading('9. PSEUDOCODE', level=2)
    pseudo = '''FUNCTION handle_chat_request(request):
    user_msg = request.getMessage()
    history = request.getHistory()
    
    IF user_msg IS empty:
        RETURN Error("Message required")
        
    system_prompt = "You are an Admission Assistant. Use this data: " + load(college_data)
    
    api_messages = [ {role: "system", content: system_prompt} ]
    FOR message IN history:
        api_messages.append(message)
    api_messages.append({role: "user", content: user_msg})
    
    TRY:
        response = call_llm_api(messages=api_messages)
        RETURN Success(response.text)
    CATCH Error as e:
        RETURN Error("Failed to connect: " + e.message)'''
    doc.add_paragraph(pseudo, style='No Spacing')

    # 10. Technologies Used
    doc.add_heading('10. TECHNOLOGIES USED', level=2)
    doc.add_paragraph('• Frontend: HTML5, CSS3, JavaScript', style='List Bullet')
    doc.add_paragraph('• Backend: Python (Flask Framework)', style='List Bullet')
    doc.add_paragraph('• LLM Integration: OpenAI Python SDK (openai==0.27.8)', style='List Bullet')
    doc.add_paragraph('• Configuration: python-dotenv (.env variables)', style='List Bullet')

    # 11. Implementation
    doc.add_heading('11. IMPLEMENTATION', level=2)
    doc.add_paragraph('The backend (app.py) handles rendering the HTML template and exposes a /chat endpoint. The frontend uses the JavaScript Fetch API to send JSON requests containing the user\'s message and an array of the last 5 messages to maintain context. The OpenAI API processes the combined prompt and context, generating responses restricted by the system rules. Error handling ensures the UI informs the user of missing API keys or network timeouts.')

    # 12. Source Code
    doc.add_heading('12. SOURCE CODE', level=2)
    doc.add_paragraph('The complete project includes app.py, templates/index.html, static/script.js, static/style.css, and data/college_data.json. Example snippet of API call in app.py:')
    code = '''response = openai.ChatCompletion.create(
    model="gpt-3.5-turbo",
    messages=messages,
    temperature=0.3,
    max_tokens=300
)'''
    doc.add_paragraph(code, style='No Spacing')

    # 13. Testing
    doc.add_heading('13. TESTING', level=2)
    doc.add_paragraph('1. Courses: "What courses are available?"', style='List Bullet')
    doc.add_paragraph('2. Eligibility: "What are the eligibility requirements for CSE?"', style='List Bullet')
    doc.add_paragraph('3. Follow-up Context: "What is the fee for the second one?"', style='List Bullet')
    doc.add_paragraph('4. Out-of-domain: "Who is the Prime Minister?"', style='List Bullet')

    # 14. Sample Output
    doc.add_heading('14. SAMPLE OUTPUT', level=2)
    doc.add_paragraph('User: What courses are available?\n'
                      'Assistant: ABC Institute of Technology offers the following courses:\n'
                      '• B.Tech Computer Science and Engineering (CSE)\n'
                      '• B.Tech Artificial Intelligence and Machine Learning (AI & ML)\n'
                      '• B.Tech Electronics and Communication Engineering (ECE)\n'
                      '• B.Tech Mechanical Engineering\n')
    doc.add_paragraph('User: Who is the Prime Minister?\n'
                      'Assistant: I apologize, but I specialize only in admission queries for ABC Institute of Technology. I am unable to answer that question.')
    
    doc.add_paragraph('[Placeholders for UI Screenshots: Dashboard, Contextual Chat, Error State]')

    # 15. Result
    doc.add_heading('15. RESULT', level=2)
    doc.add_paragraph('The LLM-powered conversational agent was successfully implemented. It accurately answers college admission queries using exclusively the provided synthetic knowledge base. Contextual continuity was successfully maintained for follow-up questions, and the agent robustly rejected out-of-domain prompts.')

    # 16. Conclusion
    doc.add_heading('16. CONCLUSION', level=2)
    doc.add_paragraph('This experiment demonstrated how Large Language Models can be programmatically constrained using prompt engineering to act as specialized domain assistants. Integrating the API through a backend ensures security (hiding API keys) and control over the injected context, allowing for a seamless conversational user experience.')

    # 17. Applications
    doc.add_heading('17. APPLICATIONS', level=2)
    doc.add_paragraph('• University Admission & Student Helpdesks', style='List Bullet')
    doc.add_paragraph('• Enterprise Customer Support & Ticketing', style='List Bullet')
    doc.add_paragraph('• Healthcare Triage Assistants', style='List Bullet')
    doc.add_paragraph('• E-Commerce Shopping Guides', style='List Bullet')

    doc.save('College_Admission_LLM_Chatbot_Experiment.docx')
    print("Report generated successfully.")

if __name__ == "__main__":
    generate_word_document()
